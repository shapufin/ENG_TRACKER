"""Regression tests for report leave day aggregation and exports."""

from datetime import date
import io

import openpyxl
from django.contrib.auth.models import User
from django.http import StreamingHttpResponse
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.db import connection
from rest_framework.test import APIClient
from rest_framework import status

from apps.leave_management.models import LeaveRequest, count_business_days
from apps.overtime.models import Client, OvertimeLog
from apps.users.models import Team, TeamMembership
from apps.reports.models.core import AuditLog
from apps.reports.services.export_service import export_service


def _streaming_body(response) -> str:
    """Materialize a StreamingHttpResponse body as text.

    The Django test client does not buffer streaming responses into
    ``response.content``; callers must join ``streaming_content`` manually.
    """
    if isinstance(response, StreamingHttpResponse):
        return b"".join(response.streaming_content).decode("utf-8")
    return response.content.decode("utf-8")


class ReportLeaveBusinessDaysTests(TestCase):
    """Detailed/Excel leave totals must use business days, not calendar span."""

    def setUp(self):
        self.client = APIClient()
        self.staff = User.objects.create_user(
            username='report-staff', password='testpass', is_staff=True,
        )
        self.employee = User.objects.create_user(
            username='report-emp', password='testpass',
        )
        # Fri–Mon range: 4 calendar days, 2 business days.
        self.start = date(2026, 7, 24)
        self.end = date(2026, 7, 27)
        self.expected = count_business_days(self.start, self.end)
        self.assertEqual(self.expected, 2)

        LeaveRequest.objects.create(
            user=self.employee,
            request_type='vacation',
            start_date=self.start,
            end_date=self.end,
            status='approved',
            reason='weekend-spanning',
        )

    def test_summary_report_uses_business_days(self):
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/summary/',
            {
                'start_date': '2026-07-01',
                'end_date': '2026-07-31',
                'report_type': 'leave',
            },
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        leave = response.data.get('leave') or response.data
        total_days = leave.get('total_days') if isinstance(leave, dict) else None
        if total_days is None and isinstance(response.data, dict):
            total_days = response.data.get('total_days')
        self.assertEqual(total_days, self.expected)

    def test_detailed_user_report_uses_business_days(self):
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/detailed/',
            {
                'start_date': '2026-07-01',
                'end_date': '2026-07-31',
                'report_type': 'leave',
                'group_by': 'user',
            },
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        users = response.data.get('users', [])
        self.assertEqual(len(users), 1)
        self.assertEqual(users[0]['leave']['total_days'], self.expected)
        self.assertEqual(users[0]['leave']['approved_days'], self.expected)

    def test_detailed_month_report_uses_business_days(self):
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/detailed/',
            {
                'start_date': '2026-07-01',
                'end_date': '2026-07-31',
                'report_type': 'leave',
                'group_by': 'month',
            },
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        leave_rows = response.data.get('leave', [])
        self.assertEqual(len(leave_rows), 1)
        self.assertEqual(leave_rows[0]['total_days'], self.expected)


class TopTeamLeadersQueryTests(TestCase):
    """Primary-team resolution must be deterministic and query-bounded."""

    def setUp(self):
        self.client = APIClient()
        self.staff = User.objects.create_user(
            username='ranking-staff', password='testpass', is_staff=True,
        )
        self.tl = User.objects.create_user(username='ranking-tl', password='testpass')
        self.member = User.objects.create_user(username='ranking-member', password='testpass')
        self.member.profile.italian_tl = self.tl
        self.member.profile.save(update_fields=['italian_tl'])
        self.client_model = Client.objects.create(name='Ranking Client', code='RANK')
        self.team_a = Team.objects.create(name='Primary Team', code='PRI', team_leader=self.tl)
        self.team_b = Team.objects.create(name='Secondary Team', code='SEC')
        TeamMembership.objects.create(
            user_profile=self.tl.profile, team=self.team_a, is_primary_team=True,
        )
        TeamMembership.objects.create(
            user_profile=self.tl.profile, team=self.team_b, is_primary_team=False,
        )
        OvertimeLog.objects.create(
            user=self.member,
            client=self.client_model,
            date=date(2026, 7, 20),
            hours=4,
            status='approved',
        )

    def test_top_team_leaders_returns_primary_team(self):
        self.client.force_authenticate(user=self.staff)
        response = self.client.get('/api/reports/top-team-leaders/')
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data[0]['team_name'], self.team_a.name)

    def test_top_team_leaders_query_count_does_not_scale_per_ranked_user(self):
        self.client.force_authenticate(user=self.staff)
        with CaptureQueriesContext(connection) as queries:
            response = self.client.get('/api/reports/top-team-leaders/')
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertLessEqual(len(queries), 12)


class LeaveExportTests(TestCase):
    """Leave export must not 500 on valid LeaveRequest fields.

    Regression: export_leave_to_excel referenced ``leave_type`` and
    ``description`` which do not exist on ``LeaveRequest`` (actual fields are
    ``request_type`` and ``reason``). The AttributeError was caught and
    returned as a 500.
    """

    def setUp(self):
        self.client = APIClient()
        self.staff = User.objects.create_user(
            username='export-staff', password='testpass', is_staff=True,
        )
        self.employee = User.objects.create_user(
            username='export-emp', password='testpass',
            first_name='Export', last_name='User',
        )
        LeaveRequest.objects.create(
            user=self.employee,
            request_type='vacation',
            start_date=date(2026, 5, 1),
            end_date=date(2026, 5, 5),
            status='approved',
            reason='Summer break',
        )

    def test_export_leave_endpoint_returns_200(self):
        """The /reports/export-leave/ endpoint must not 500."""
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/export-leave/',
            {'start_date': '2026-05-01', 'end_date': '2026-07-01'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def test_export_leave_service_produces_valid_xlsx(self):
        """The export service must produce a readable .xlsx byte stream."""
        leave_qs = LeaveRequest.objects.filter(status='approved')
        result = export_service.export_leave_to_excel(leave_qs)
        self.assertIsInstance(result, bytes)
        self.assertGreater(len(result), 0)
        # Verify it's a valid zip-based xlsx (PK header)
        self.assertEqual(result[:2], b'PK')


class AuditLogExportTests(TestCase):
    """The /reports/audit-logs/export/ endpoint must return CSV.

    Regression: the route previously 404'd because AuditLogViewSet had no
    ``export`` action; DRF resolved ``/export/`` as a retrieve with pk="export".
    """

    def setUp(self):
        self.client = APIClient()
        self.staff = User.objects.create_user(
            username='audit-staff', password='testpass', is_staff=True,
        )
        self.employee = User.objects.create_user(
            username='audit-emp', password='testpass',
        )
        self.other = User.objects.create_user(
            username='audit-other', password='testpass',
        )
        self.log_emp = AuditLog.objects.create(
            user=self.employee,
            action='CREATE',
            model_name='overtimelog',
            object_id='1',
            object_repr='OT #1',
            new_values={'hours': 4},
        )
        self.log_other = AuditLog.objects.create(
            user=self.other,
            action='DELETE',
            model_name='leaverequest',
            object_id='2',
            object_repr='Leave #2',
        )

    def test_export_returns_csv(self):
        self.client.force_authenticate(user=self.staff)
        response = self.client.get('/api/reports/audit-logs/export/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'text/csv')
        body = _streaming_body(response)
        self.assertIn('User', body.split('\r\n')[0])
        self.assertIn('audit-emp', body)
        self.assertIn('audit-other', body)

    def test_export_respects_action_filter(self):
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/audit-logs/export/', {'action': 'DELETE'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        body = _streaming_body(response)
        self.assertIn('audit-other', body)
        self.assertNotIn('audit-emp', body)

    def test_non_staff_sees_only_own_logs(self):
        self.client.force_authenticate(user=self.employee)
        response = self.client.get('/api/reports/audit-logs/export/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        body = _streaming_body(response)
        self.assertIn('audit-emp', body)
        self.assertNotIn('audit-other', body)

    def test_export_streams_large_dataset(self):
        """StreamingHttpResponse must yield every row, not truncate.

        Regression guard: the previous ``HttpResponse`` implementation
        buffered the entire body in memory; the streaming generator must
        yield all rows without dropping any at chunk boundaries.
        """
        from django.utils import timezone
        # Bulk-create 1500 audit logs (well above the 1000-row chunk_size
        # used by ``queryset.iterator(chunk_size=1000)``).
        logs = [
            AuditLog(
                user=self.employee,
                action='UPDATE',
                model_name='overtimelog',
                object_id=str(i),
                object_repr=f'OT #{i}',
                timestamp=timezone.now(),
            )
            for i in range(1500)
        ]
        AuditLog.objects.bulk_create(logs)
        self.client.force_authenticate(user=self.staff)
        response = self.client.get('/api/reports/audit-logs/export/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        body = _streaming_body(response)
        # Header + 1500 bulk-created + 2 setUp logs = 1503 lines.
        lines = [ln for ln in body.split('\r\n') if ln]
        self.assertEqual(len(lines), 1503, 'expected header + 1502 data rows')


class AuditLogStatsTests(TestCase):
    """The /reports/audit-logs/stats/ endpoint aggregates over the full
    filtered queryset, not the page-limited ``results`` array.

    Regression: ``useAuditLogs`` previously computed stats client-side
    from ``logsData.results`` (page_size=100). The stats cards therefore
    reflected only the first 100 rows of the filtered set.
    """

    def setUp(self):
        self.client = APIClient()
        self.staff = User.objects.create_user(
            username='stats-staff', password='testpass', is_staff=True,
        )
        self.user_a = User.objects.create_user(
            username='stats-a', password='testpass',
        )
        self.user_b = User.objects.create_user(
            username='stats-b', password='testpass',
        )
        self.user_c = User.objects.create_user(
            username='stats-c', password='testpass',
        )

    def _create_logs(self, user, action, count, **extra):
        from django.utils import timezone
        defaults = dict(
            model_name='overtimelog',
            object_id='1',
            object_repr='OT #1',
            timestamp=timezone.now(),
        )
        defaults.update(extra)
        AuditLog.objects.bulk_create([
            AuditLog(user=user, action=action, **defaults) for _ in range(count)
        ])

    def test_stats_returns_full_filtered_set(self):
        """200 logs → total_logs == 200 (not capped at 100)."""
        self._create_logs(self.user_a, 'CREATE', 200)
        self.client.force_authenticate(user=self.staff)
        response = self.client.get('/api/reports/audit-logs/stats/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['total_logs'], 200)

    def test_stats_respects_action_filter(self):
        self._create_logs(self.user_a, 'DELETE', 5)
        self._create_logs(self.user_a, 'CREATE', 10)
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/audit-logs/stats/', {'action': 'DELETE'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['total_logs'], 5)
        self.assertEqual(response.data['failed_actions'], 5)

    def test_stats_unique_users(self):
        self._create_logs(self.user_a, 'CREATE', 10)
        self._create_logs(self.user_b, 'CREATE', 10)
        self._create_logs(self.user_c, 'CREATE', 10)
        self.client.force_authenticate(user=self.staff)
        response = self.client.get('/api/reports/audit-logs/stats/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['unique_users'], 3)

    def test_stats_success_rate(self):
        """100 total, 5 DELETE/REJECT → success_rate == 95.0."""
        self._create_logs(self.user_a, 'CREATE', 95)
        self._create_logs(self.user_a, 'DELETE', 5)
        self.client.force_authenticate(user=self.staff)
        response = self.client.get('/api/reports/audit-logs/stats/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['total_logs'], 100)
        self.assertEqual(response.data['failed_actions'], 5)
        self.assertEqual(response.data['success_rate'], 95.0)

    def test_stats_empty_queryset(self):
        """No logs match filter → total_logs == 0, success_rate == 100."""
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/audit-logs/stats/', {'action': 'DELETE'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['total_logs'], 0)
        self.assertEqual(response.data['success_rate'], 100)

    def test_stats_counts_bulk_reject_as_failed(self):
        """BULK_REJECT must be counted as a failed action (not just DELETE/REJECT)."""
        self._create_logs(self.user_a, 'CREATE', 90)
        self._create_logs(self.user_a, 'REJECT', 3)
        self._create_logs(self.user_a, 'BULK_REJECT', 7)
        self.client.force_authenticate(user=self.staff)
        response = self.client.get('/api/reports/audit-logs/stats/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['total_logs'], 100)
        self.assertEqual(response.data['failed_actions'], 10)  # 3 REJECT + 7 BULK_REJECT


class MixedTeamLeaderExportTests(TestCase):
    """Exports must not crash when TL groups mix None and integer IDs.

    Regression: ``sorted(tl_groups.items())`` raised TypeError on Python 3
    when keys contained both ``None`` (unassigned) and integer TL IDs.
    """

    def setUp(self):
        self.tl = User.objects.create_user(username='mix-tl', password='testpass')
        self.assigned = User.objects.create_user(
            username='mix-assigned', password='testpass',
        )
        self.unassigned = User.objects.create_user(
            username='mix-unassigned', password='testpass',
        )
        self.assigned.profile.italian_tl = self.tl
        self.assigned.profile.save(update_fields=['italian_tl'])
        self.client_model = Client.objects.create(name='Mix Client', code='MIX')
        OvertimeLog.objects.create(
            user=self.assigned,
            client=self.client_model,
            date=date(2026, 6, 1),
            hours=4,
            status='approved',
        )
        OvertimeLog.objects.create(
            user=self.unassigned,
            client=self.client_model,
            date=date(2026, 6, 2),
            hours=3,
            status='approved',
        )
        LeaveRequest.objects.create(
            user=self.assigned,
            request_type='vacation',
            start_date=date(2026, 6, 10),
            end_date=date(2026, 6, 12),
            status='approved',
            reason='mix',
        )
        LeaveRequest.objects.create(
            user=self.unassigned,
            request_type='vacation',
            start_date=date(2026, 6, 15),
            end_date=date(2026, 6, 17),
            status='approved',
            reason='mix',
        )

    def test_ot_standby_export_with_mixed_tl_assignment(self):
        ot_qs = OvertimeLog.objects.filter(status='approved')
        sb_qs = OvertimeLog.objects.none()
        result = export_service.export_ot_standby_to_excel(ot_qs, sb_qs)
        self.assertIsInstance(result, bytes)
        self.assertEqual(result[:2], b'PK')


class ExportPendingStatusTests(TestCase):
    """The /reports/export-ot-standby/?status=pending endpoint must export
    pending records (not approved) so HR can review the pending backlog."""

    def setUp(self):
        self.client = APIClient()
        self.staff = User.objects.create_user(
            username='pending-staff', password='testpass', is_staff=True,
        )
        self.employee = User.objects.create_user(
            username='pending-emp', password='testpass',
            first_name='Pending', last_name='Worker',
        )
        self.client_model = Client.objects.create(name='Pending Client', code='PND')
        # One approved + one pending OT record on the same date.
        OvertimeLog.objects.create(
            user=self.employee, client=self.client_model,
            date=date(2026, 7, 10), hours=4, status='approved',
        )
        OvertimeLog.objects.create(
            user=self.employee, client=self.client_model,
            date=date(2026, 7, 11), hours=6, status='pending',
        )

    def test_default_export_returns_approved_only(self):
        """Default (no status param) exports approved records only."""
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/export-ot-standby/',
            {'start_date': '2026-07-01', 'end_date': '2026-07-31'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        # Find the Status column (L) values in data rows.
        statuses = [ws.cell(row=r, column=12).value for r in range(2, ws.max_row + 1)]
        statuses = [s for s in statuses if s]
        self.assertTrue(all(s == 'Approved' for s in statuses), f'expected only Approved, got {statuses}')

    def test_pending_export_returns_pending_only(self):
        """status=pending exports pending records only."""
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/export-ot-standby/',
            {'start_date': '2026-07-01', 'end_date': '2026-07-31', 'status': 'pending'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        statuses = [ws.cell(row=r, column=12).value for r in range(2, ws.max_row + 1)]
        statuses = [s for s in statuses if s]
        self.assertTrue(statuses, 'expected at least one pending row')
        self.assertTrue(all(s == 'Pending' for s in statuses), f'expected only Pending, got {statuses}')
        # Filename should include _pending suffix.
        self.assertIn('_pending', response['Content-Disposition'])

    def test_leave_export_with_mixed_tl_assignment(self):
        leave_qs = LeaveRequest.objects.filter(status='approved')
        result = export_service.export_leave_to_excel(leave_qs)
        self.assertIsInstance(result, bytes)
        self.assertEqual(result[:2], b'PK')


class ExportGroupingAndStylingTests(TestCase):
    """OT/Standby and Leave exports must group records by user (not interleave
    them across dates) and color-code the OT/Standby Type column.

    Regression: records were sorted ``(date, username)``, so a given user's
    records were scattered across the sheet. The Type column showed
    "Extra hours"/"On Call Service" with no visual distinction.
    """

    def setUp(self):
        import openpyxl
        self.openpyxl = openpyxl
        # Two users with interleaved dates so date-first sort would mix them.
        self.alpha = User.objects.create_user(
            username='alpha-emp', password='testpass',
            first_name='Alpha', last_name='Employee',
        )
        self.beta = User.objects.create_user(
            username='beta-emp', password='testpass',
            first_name='Beta', last_name='Employee',
        )
        self.client_model = Client.objects.create(name='Grouping Client', code='GRP')
        # Alpha: OT on July 6, SB on July 8, OT on July 9.
        # Beta: OT on July 7.
        # Date-first order would interleave both users AND categories:
        #   Alpha(7/6 OT), Beta(7/7 OT), Alpha(7/8 SB), Alpha(7/9 OT).
        # User+category order must be:
        #   Alpha(7/6 OT), Alpha(7/9 OT), Alpha(7/8 SB), Beta(7/7 OT).
        OvertimeLog.objects.create(
            user=self.alpha, client=self.client_model,
            date=date(2026, 7, 6), hours=4, status='approved',
        )
        OvertimeLog.objects.create(
            user=self.beta, client=self.client_model,
            date=date(2026, 7, 7), hours=3, status='approved',
        )
        from apps.standby.models import StandbyLog
        StandbyLog.objects.create(
            user=self.alpha, date=date(2026, 7, 8), hours=2, status='approved',
        )
        OvertimeLog.objects.create(
            user=self.alpha, client=self.client_model,
            date=date(2026, 7, 9), hours=5, status='approved',
        )
        # Leave: Alpha starts July 10, Beta starts July 9 — date-first would
        # put Beta before Alpha; user-first must keep Alpha before Beta.
        LeaveRequest.objects.create(
            user=self.alpha, request_type='vacation',
            start_date=date(2026, 7, 10), end_date=date(2026, 7, 10),
            status='approved', reason='a',
        )
        LeaveRequest.objects.create(
            user=self.beta, request_type='vacation',
            start_date=date(2026, 7, 9), end_date=date(2026, 7, 9),
            status='approved', reason='b',
        )

    def _load_wb(self, data):
        import io
        return self.openpyxl.load_workbook(io.BytesIO(data))

    def test_ot_standby_records_grouped_by_user_then_category(self):
        """Within a TL sheet, each user's records are contiguous AND within a
        user the categories are contiguous (all Extra hours, then all On Call
        Service) — sorted by (user, type, date)."""
        ot_qs = OvertimeLog.objects.filter(status='approved')
        from apps.standby.models import StandbyLog
        sb_qs = StandbyLog.objects.filter(status='approved')
        wb = self._load_wb(export_service.export_ot_standby_to_excel(ot_qs, sb_qs))
        ws = wb.active
        # Collect (name, type) pairs in row order.
        rows = [
            (ws.cell(row=r, column=1).value, ws.cell(row=r, column=7).value)
            for r in range(2, ws.max_row + 1)
        ]
        rows = [(n, t) for n, t in rows if n]
        names = [n for n, _ in rows]
        # Alpha has 3 records (2 OT + 1 SB), Beta has 1 (OT).
        alpha_indices = [i for i, n in enumerate(names) if n.startswith('Alpha')]
        beta_indices = [i for i, n in enumerate(names) if n.startswith('Beta')]
        self.assertEqual(len(alpha_indices), 3)
        self.assertEqual(len(beta_indices), 1)
        # Alpha's 3 rows must be contiguous (grouped), not separated by Beta.
        self.assertEqual(alpha_indices[-1] - alpha_indices[0], 2)
        # Beta comes after Alpha (alphabetical username order).
        self.assertGreater(beta_indices[0], alpha_indices[-1])
        # Within Alpha, both OT (Extra hours) rows must come before the SB
        # (On Call Service) row — categories contiguous, not interleaved.
        alpha_types = [t for n, t in rows if n.startswith('Alpha')]
        self.assertEqual(
            alpha_types,
            ['Extra hours', 'Extra hours', 'On Call Service'],
            'Alpha categories must be grouped: all OT, then all SB',
        )
        # And the OT dates within Alpha must be ascending (7/6 before 7/9).
        alpha_ot_dates = [
            ws.cell(row=r, column=2).value
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=1).value
            and ws.cell(row=r, column=1).value.startswith('Alpha')
            and ws.cell(row=r, column=7).value == 'Extra hours'
        ]
        self.assertEqual(alpha_ot_dates, ['06/07/2026', '09/07/2026'])

    def test_ot_standby_type_cells_color_coded(self):
        """The Type column (G) has a distinct fill per category."""
        from apps.reports.services.export_service import OT_TYPE_COLOR, SB_TYPE_COLOR
        ot_qs = OvertimeLog.objects.filter(status='approved')
        from apps.standby.models import StandbyLog
        sb_qs = StandbyLog.objects.filter(status='approved')
        wb = self._load_wb(export_service.export_ot_standby_to_excel(ot_qs, sb_qs))
        ws = wb.active
        ot_fills, sb_fills = set(), set()
        for r in range(2, ws.max_row + 1):
            type_cell = ws.cell(row=r, column=7)
            fill = type_cell.fill.fgColor.rgb if type_cell.fill and type_cell.fill.fgColor else None
            if type_cell.value == 'Extra hours':
                ot_fills.add(fill)
            elif type_cell.value == 'On Call Service':
                sb_fills.add(fill)
        # Both categories present with non-default fills.
        self.assertTrue(ot_fills, 'no Extra hours rows found')
        self.assertTrue(sb_fills, 'no On Call Service rows found')
        # openpyxl returns ARGB like '00BDD7EE' for solid fills; check the
        # configured hex appears and the two categories differ.
        self.assertTrue(any(OT_TYPE_COLOR in (f or '') for f in ot_fills))
        self.assertTrue(any(SB_TYPE_COLOR in (f or '') for f in sb_fills))
        self.assertNotEqual(ot_fills, sb_fills)

    def test_leave_records_grouped_by_user(self):
        """Leave records are sorted by user, then start date."""
        leave_qs = LeaveRequest.objects.filter(status='approved')
        wb = self._load_wb(export_service.export_leave_to_excel(leave_qs))
        ws = wb.active
        names = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        names = [n for n in names if n]
        # Alpha (username alpha-emp) must come before Beta despite Beta's
        # earlier start date.
        self.assertEqual(names, ['Alpha Employee', 'Beta Employee'])


class ExportVisibilityTests(TestCase):
    """Export endpoints must scope querysets by the caller's visibility.

    Regression: ExportOTStandbyView and ExportLeaveView built their own
    querysets without applying _apply_visibility_constraints (unlike
    UnifiedReportFilter used by ExportExcelView). A team leader or regular
    employee could export ALL records, not just their own/team's.
    """

    def setUp(self):
        from apps.users.models import UserProfile
        self.client = APIClient()
        self.staff = User.objects.create_user(
            username='vis-staff', password='testpass', is_staff=True,
        )
        self.tl = User.objects.create_user(
            username='vis-tl', password='testpass',
            first_name='Team', last_name='Leader',
        )
        # Profile auto-created via signal; update it.
        self.tl_profile = UserProfile.objects.get(user=self.tl)
        self.tl_profile.is_italian_tl_role = True
        self.tl_profile.save(update_fields=['is_italian_tl_role'])
        self.emp_a = User.objects.create_user(
            username='vis-emp-a', password='testpass',
            first_name='EmpA', last_name='User',
        )
        self.emp_b = User.objects.create_user(
            username='vis-emp-b', password='testpass',
            first_name='EmpB', last_name='User',
        )
        prof_a = UserProfile.objects.get(user=self.emp_a)
        prof_a.italian_tl = self.tl
        prof_a.save(update_fields=['italian_tl'])
        # emp_b profile: no TL assignment (NOT on TL's team)
        self.client_model = Client.objects.create(name='Vis Client', code='VIS')
        # OT for emp_a (on TL's team) and emp_b (NOT on TL's team)
        OvertimeLog.objects.create(
            user=self.emp_a, client=self.client_model,
            date=date(2026, 7, 10), hours=4, status='approved',
        )
        OvertimeLog.objects.create(
            user=self.emp_b, client=self.client_model,
            date=date(2026, 7, 11), hours=6, status='approved',
        )
        # Leave for emp_a and emp_b
        LeaveRequest.objects.create(
            user=self.emp_a, request_type='vacation',
            start_date=date(2026, 7, 1), end_date=date(2026, 7, 3),
            status='approved', reason='A leave',
        )
        LeaveRequest.objects.create(
            user=self.emp_b, request_type='vacation',
            start_date=date(2026, 7, 5), end_date=date(2026, 7, 7),
            status='approved', reason='B leave',
        )

    def test_tl_export_ot_standby_excludes_non_team(self):
        """TL must not see emp_b's records in OT/standby export."""
        self.client.force_authenticate(user=self.tl)
        response = self.client.get(
            '/api/reports/export-ot-standby/',
            {'start_date': '2026-07-01', 'end_date': '2026-07-31'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        # Column 1 = Name (full name), column 2 = Date
        names = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        names = [n for n in names if n]
        self.assertIn('EmpA User', names)
        self.assertNotIn('EmpB User', names,
                         'TL export leaked non-team member records')

    def test_tl_export_leave_excludes_non_team(self):
        """TL must not see emp_b's records in leave export."""
        self.client.force_authenticate(user=self.tl)
        response = self.client.get(
            '/api/reports/export-leave/',
            {'start_date': '2026-07-01', 'end_date': '2026-07-31'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        # Leave export columns: User (col 1), Username (col 2) — check both
        names = [ws.cell(row=r, column=c).value
                 for r in range(2, ws.max_row + 1)
                 for c in range(1, 3)]
        names_str = ' '.join(str(n) for n in names if n)
        self.assertIn('EmpA', names_str)
        self.assertNotIn('EmpB', names_str,
                         'TL leave export leaked non-team member records')

    def test_employee_export_ot_standby_only_own(self):
        """Regular employee must only see their own records."""
        self.client.force_authenticate(user=self.emp_a)
        response = self.client.get(
            '/api/reports/export-ot-standby/',
            {'start_date': '2026-07-01', 'end_date': '2026-07-31'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        names = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        names = [n for n in names if n]
        self.assertEqual(names, ['EmpA User'])

    def test_export_populates_evidence_from_ticket_references(self):
        """Ticket-type OT entries populate Evidence + Reference Code from ticket_references."""
        OvertimeLog.objects.create(
            user=self.emp_a, client=self.client_model,
            date=date(2026, 7, 12), hours=3, status='approved',
            evidence_type='ticket',
            ticket_references=['INC-100', 'INC-200'],
        )
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/export-ot-standby/',
            {'start_date': '2026-07-01', 'end_date': '2026-07-31'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        # Export groups by Italian TL into separate sheets; check all sheets
        all_evidence = []
        all_ref_codes = []
        for ws in wb.worksheets:
            for r in range(2, ws.max_row + 1):
                all_evidence.append(ws.cell(row=r, column=10).value)
                all_ref_codes.append(ws.cell(row=r, column=11).value)
        self.assertIn('INC-100, INC-200', all_evidence)
        self.assertIn('INC-100, INC-200', all_ref_codes)

    def test_export_preserves_existing_ticket_values(self):
        """Structured refs do not overwrite existing evidence or reference code."""
        OvertimeLog.objects.create(
            user=self.emp_a, client=self.client_model,
            date=date(2026, 7, 13), hours=2, status='approved',
            evidence_type='ticket',
            evidence='Legacy ticket context',
            reference_code='Legacy-REF',
            ticket_references=['INC-300'],
        )
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/export-ot-standby/',
            {'start_date': '2026-07-01', 'end_date': '2026-07-31'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        all_values = [
            ws.cell(row=r, column=c).value
            for ws in wb.worksheets
            for r in range(2, ws.max_row + 1)
            for c in (10, 11)
        ]
        self.assertIn('Legacy ticket context', all_values)
        self.assertIn('Legacy-REF', all_values)


class PayrollPeriodExportTests(TestCase):
    """date_mode=processing_period exports OT/standby by requested_processing_period
    (not work date) and includes carryover columns. Used for HR-to-Payroll reports
    so carried-over entries appear in the settlement month, not the work-date month."""

    def setUp(self):
        self.client = APIClient()
        self.staff = User.objects.create_user(
            username='payroll-staff', password='testpass', is_staff=True,
        )
        self.employee = User.objects.create_user(
            username='payroll-emp', password='testpass',
            first_name='Payroll', last_name='Worker',
        )
        self.client_model = Client.objects.create(name='Payroll Client', code='PAY')
        # Regular May entry (work date May, processing period May)
        OvertimeLog.objects.create(
            user=self.employee, client=self.client_model,
            date=date(2026, 5, 15), hours=4, status='approved',
        )
        # Carried-over May entry (work date May, processing period June)
        carried = OvertimeLog.objects.create(
            user=self.employee, client=self.client_model,
            date=date(2026, 5, 30), hours=2, status='approved',
        )
        carried.requested_processing_period = date(2026, 6, 1)
        carried.save(update_fields=['requested_processing_period'])

    def test_work_date_mode_excludes_carryover_from_june(self):
        """Default (work_date) June export excludes May-dated carryover."""
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/export-ot-standby/',
            {'start_date': '2026-06-01', 'end_date': '2026-06-30'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        dates = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
        dates = [d for d in dates if d]
        self.assertEqual(dates, [], 'work-date June export should have no rows')

    def test_processing_period_mode_includes_carryover_in_june(self):
        """processing_period June export includes May-dated carryover entries."""
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/export-ot-standby/',
            {'start_date': '2026-06-01', 'end_date': '2026-06-30',
             'date_mode': 'processing_period'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        # Column B = work date, Column M = processing period, Column N = carried over
        dates = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
        dates = [d for d in dates if d]
        self.assertEqual(len(dates), 1, 'should have exactly the carryover row')
        # Work date is May 30 (30/05/2026)
        self.assertIn('30/05/2026', dates[0])
        # Processing period column (M=13) should be 06/2026
        pp = ws.cell(row=2, column=13).value
        self.assertEqual(pp, '06/2026')
        # Carried over column (N=14) should be 'Yes'
        carried = ws.cell(row=2, column=14).value
        self.assertEqual(carried, 'Yes')
        # Filename should include _payroll suffix
        self.assertIn('_payroll', response['Content-Disposition'])

    def test_processing_period_mode_excludes_carryover_from_may(self):
        """processing_period May export excludes entries carried to June."""
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/export-ot-standby/',
            {'start_date': '2026-05-01', 'end_date': '2026-05-31',
             'date_mode': 'processing_period'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        dates = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
        dates = [d for d in dates if d]
        # Only the regular May entry (15/05/2026), not the carryover (30/05 carried to June)
        self.assertEqual(len(dates), 1)
        self.assertIn('15/05/2026', dates[0])

    def test_work_date_mode_also_has_carryover_columns(self):
        """Default (work_date) export also includes carryover columns."""
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(
            '/api/reports/export-ot-standby/',
            {'start_date': '2026-05-01', 'end_date': '2026-05-31'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        # Column M (13) header should be 'Processing Period'
        header_k = ws.cell(row=1, column=13).value
        self.assertEqual(header_k, 'Processing Period')
        # Column N (14) header should be 'Carried Over'
        header_l = ws.cell(row=1, column=14).value
        self.assertEqual(header_l, 'Carried Over')

