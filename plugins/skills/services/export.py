"""Exports of the team skill matrix.

Two shapes, both reachable:

* ``export_matrix_csv`` — long format, one row per rating. The original.
* ``export_matrix_xlsx`` — wide format, one row per person and one column per
  skill, grouped under a category band row, levels written as their display
  names. Matches the spreadsheet the business already keeps by hand.
"""
import csv
import io

from django.contrib.auth.models import User
from django.db.models import Q
from django.http import HttpResponse, StreamingHttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from plugins.skills.models import SkillLevelLabels, UserSkill, Skill


class Echo:
    """An object that implements just the write method of the file-like
    interface. Used to stream CSV rows without buffering the whole file."""
    def write(self, value):
        return value


def export_matrix_csv(visible_ids, category_code=None, search=None):
    """Return a StreamingHttpResponse with the team skill matrix as CSV.

    Columns: Username, Skill, Category, Level, Notes.
    One row per UserSkill (not per user).
    """
    skills_qs = Skill.objects.filter(is_active=True)
    if category_code:
        skills_qs = skills_qs.filter(category__code=category_code)
    skill_ids = list(skills_qs.values_list('id', flat=True))

    # Apply search filter: narrow visible_ids to users matching the search term.
    if search:
        from django.contrib.auth.models import User
        from django.db.models import Q
        matching = User.objects.filter(
            is_active=True,
            id__in=visible_ids,
        ).filter(
            Q(username__icontains=search)
            | Q(first_name__icontains=search)
            | Q(last_name__icontains=search)
            | Q(user_skills__skill__name__icontains=search)
        ).distinct().values_list('id', flat=True)
        visible_ids = set(matching)

    ratings = UserSkill.objects.filter(
        user_id__in=visible_ids,
        skill_id__in=skill_ids,
    ).select_related('user', 'skill', 'skill__category').order_by(
        'user__username', 'skill__category__name', 'skill__name'
    )

    rows = [['Username', 'Skill', 'Category', 'Level', 'Notes']]
    for r in ratings:
        rows.append([
            r.user.username,
            r.skill.name,
            r.skill.category.name,
            str(r.level),
            r.notes,
        ])

    pseudo_buffer = Echo()
    writer = csv.writer(pseudo_buffer)
    response = StreamingHttpResponse(
        (writer.writerow(row) for row in rows),
        content_type="text/csv",
    )
    response['Content-Disposition'] = 'attachment; filename="skills_matrix.csv"'
    return response


XLSX_CONTENT_TYPE = (
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
)

IDENTITY_HEADERS = ['Surname', 'Name', 'Username', 'Grade']

_BAND_FILL = PatternFill('solid', fgColor='DDEBF7')
_HEADER_FILL = PatternFill('solid', fgColor='F2F2F2')


def _level_names():
    """Level number -> display name, honouring the admin-set labels.

    The singleton is the same source the UI resolves through, so a renamed
    level reads the same in the sheet as it does on screen.
    """
    labels = SkillLevelLabels.get_singleton()
    return {
        level: getattr(labels, f'level_{level}_label') for level in range(1, 6)
    }


def _search_filtered_ids(visible_ids, search):
    """Narrow ``visible_ids`` to people matching a free-text search."""
    return set(
        User.objects.filter(is_active=True, id__in=visible_ids)
        .filter(
            Q(username__icontains=search)
            | Q(first_name__icontains=search)
            | Q(last_name__icontains=search)
            | Q(user_skills__skill__name__icontains=search)
        )
        .distinct()
        .values_list('id', flat=True)
    )


def _grade_for(profile):
    """"Infrastructure L3" style summary of the person's Tech grades.

    Reads the ``tech_assignments`` prefetch — the caller must set it up.
    """
    if profile is None:
        return ''
    from apps.users.services.tech_assignments import format_assignments

    return ', '.join(format_assignments(profile))


def export_matrix_xlsx(visible_ids, category_code=None, search=None):
    """Return an ``HttpResponse`` with the wide skill matrix as XLSX.

    One row per visible active person — including people with no ratings at
    all, so gaps stay visible — and one column per active skill, ordered by
    category then skill name. Cells carry the level's display name, blank when
    unrated.

    Not streamed: openpyxl writes a zip archive, which has no row-wise
    generator form. The row count is bounded by the team the caller can see.
    """
    if search:
        visible_ids = _search_filtered_ids(visible_ids, search)

    skills_qs = Skill.objects.filter(is_active=True).select_related('category')
    if category_code:
        skills_qs = skills_qs.filter(category__code=category_code)
    skills = list(skills_qs.order_by('category__name', 'name'))

    users = list(
        User.objects.filter(is_active=True, id__in=visible_ids)
        .select_related('profile')
        .prefetch_related(
            'profile__tech_assignments__tech', 'profile__tech_assignments__level'
        )
        .order_by('last_name', 'first_name', 'username')
    )

    # One pass over the ratings for every person and skill on the sheet.
    levels = {
        (user_id, skill_id): level
        for user_id, skill_id, level in UserSkill.objects.filter(
            user_id__in=visible_ids, skill_id__in=[s.id for s in skills]
        ).values_list('user_id', 'skill_id', 'level')
    }
    level_names = _level_names()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Skill matrix'

    band_row = [None] * len(IDENTITY_HEADERS)
    header_row = list(IDENTITY_HEADERS)
    previous_category = None
    for skill in skills:
        # The band label is written once per category group, over its first
        # column, exactly like the hand-kept spreadsheet.
        band_row.append(
            skill.category.name if skill.category.name != previous_category else None
        )
        previous_category = skill.category.name
        header_row.append(skill.name)
    ws.append(band_row)
    ws.append(header_row)

    for user in users:
        row = [
            user.last_name,
            user.first_name,
            user.username,
            _grade_for(getattr(user, 'profile', None)),
        ]
        for skill in skills:
            level = levels.get((user.id, skill.id))
            row.append(level_names.get(level) if level else None)
        ws.append(row)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _BAND_FILL
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(textRotation=90, vertical='bottom')
    # Freeze the identity columns and both header rows so scrolling 100+ skill
    # columns keeps the person and the skill name on screen.
    ws.freeze_panes = ws.cell(row=3, column=len(IDENTITY_HEADERS) + 1)
    for index in range(1, len(IDENTITY_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(index)].width = 18
    for index in range(len(IDENTITY_HEADERS) + 1, len(header_row) + 1):
        ws.column_dimensions[get_column_letter(index)].width = 14

    lists = wb.create_sheet('Lists')
    lists.append(['Proficiency'])
    lists['A1'].font = Font(bold=True)
    for level in range(1, 6):
        lists.append([level_names[level]])
    lists.column_dimensions['A'].width = 18

    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response['Content-Disposition'] = 'attachment; filename="skills_matrix.xlsx"'
    return response
