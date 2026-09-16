"""Tests for the skill matrix export services (long CSV + wide XLSX)."""
import csv
import io

from django.contrib.auth.models import User
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from apps.users.models.core import Tech, TechLevel, UserTech
from plugins.skills.models import SkillCategory, Skill, SkillLevelLabels, UserSkill
from plugins.skills.services.export import export_matrix_csv, export_matrix_xlsx


class ExportMatrixCsvTest(TestCase):
    def setUp(self):
        self.cat = SkillCategory.objects.create(name="Backend", code="backend")
        self.skill1 = Skill.objects.create(
            category=self.cat, name="Python", code="python"
        )
        self.skill2 = Skill.objects.create(
            category=self.cat, name="Django", code="django"
        )
        self.user1 = User.objects.create_user(username="alice", password="x")
        self.user2 = User.objects.create_user(username="bob", password="x")

    def _parse_csv(self, response):
        """Collect all streamed chunks and parse as CSV."""
        chunks = list(response.streaming_content)
        raw = b"".join(
            c.encode('utf-8') if isinstance(c, str) else c for c in chunks
        ).decode('utf-8')
        reader = csv.reader(io.StringIO(raw))
        return list(reader)

    def test_response_headers(self):
        UserSkill.objects.create(
            user=self.user1, skill=self.skill1, level=3,
            last_updated_by=self.user1,
        )
        resp = export_matrix_csv({self.user1.id})
        self.assertEqual(resp['Content-Type'], 'text/csv')
        self.assertIn('attachment', resp['Content-Disposition'])
        self.assertIn('skills_matrix.csv', resp['Content-Disposition'])

    def test_header_row(self):
        UserSkill.objects.create(
            user=self.user1, skill=self.skill1, level=3,
            last_updated_by=self.user1,
        )
        resp = export_matrix_csv({self.user1.id})
        rows = self._parse_csv(resp)
        self.assertEqual(
            rows[0],
            ['Username', 'Skill', 'Category', 'Level', 'Notes'],
        )

    def test_one_row_per_user_skill(self):
        UserSkill.objects.create(
            user=self.user1, skill=self.skill1, level=3,
            last_updated_by=self.user1,
        )
        UserSkill.objects.create(
            user=self.user1, skill=self.skill2, level=2,
            last_updated_by=self.user1,
        )
        UserSkill.objects.create(
            user=self.user2, skill=self.skill1, level=5,
            last_updated_by=self.user2,
        )
        resp = export_matrix_csv({self.user1.id, self.user2.id})
        rows = self._parse_csv(resp)
        # 1 header + 3 data rows
        self.assertEqual(len(rows), 4)

    def test_row_content(self):
        UserSkill.objects.create(
            user=self.user1, skill=self.skill1, level=3,
            last_updated_by=self.user1, notes="learning",
        )
        resp = export_matrix_csv({self.user1.id})
        rows = self._parse_csv(resp)
        data = rows[1]
        self.assertEqual(data[0], 'alice')
        self.assertEqual(data[1], 'Python')
        self.assertEqual(data[2], 'Backend')
        self.assertEqual(data[3], '3')
        self.assertEqual(data[4], 'learning')

    def test_empty_team_returns_header_only(self):
        resp = export_matrix_csv(set())
        rows = self._parse_csv(resp)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], 'Username')

    def test_category_filter(self):
        other_cat = SkillCategory.objects.create(name="Frontend", code="frontend")
        other_skill = Skill.objects.create(
            category=other_cat, name="React", code="react"
        )
        UserSkill.objects.create(
            user=self.user1, skill=self.skill1, level=3,
            last_updated_by=self.user1,
        )
        UserSkill.objects.create(
            user=self.user1, skill=other_skill, level=3,
            last_updated_by=self.user1,
        )
        resp = export_matrix_csv({self.user1.id}, category_code='backend')
        rows = self._parse_csv(resp)
        # header + 1 row (Python only; React filtered out)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1], 'Python')

    def test_inactive_skills_excluded(self):
        inactive = Skill.objects.create(
            category=self.cat, name="Legacy", code="legacy", is_active=False
        )
        UserSkill.objects.create(
            user=self.user1, skill=inactive, level=3,
            last_updated_by=self.user1,
        )
        resp = export_matrix_csv({self.user1.id})
        rows = self._parse_csv(resp)
        # header only — inactive skill excluded
        self.assertEqual(len(rows), 1)

    def test_search_filters_by_username(self):
        """Export with search= only includes matching users."""
        UserSkill.objects.create(
            user=self.user1, skill=self.skill1, level=3,
            last_updated_by=self.user1,
        )
        UserSkill.objects.create(
            user=self.user2, skill=self.skill1, level=4,
            last_updated_by=self.user2,
        )
        resp = export_matrix_csv({self.user1.id, self.user2.id}, search='alice')
        rows = self._parse_csv(resp)
        usernames = [r[0] for r in rows[1:]]  # skip header
        self.assertIn('alice', usernames)
        self.assertNotIn('bob', usernames)

    def test_search_filters_by_skill_name(self):
        """Export with search= includes users who have the matching skill."""
        UserSkill.objects.create(
            user=self.user1, skill=self.skill1, level=3,  # Python
            last_updated_by=self.user1,
        )
        UserSkill.objects.create(
            user=self.user2, skill=self.skill2, level=4,  # Django
            last_updated_by=self.user2,
        )
        resp = export_matrix_csv({self.user1.id, self.user2.id}, search='python')
        rows = self._parse_csv(resp)
        usernames = [r[0] for r in rows[1:]]
        self.assertIn('alice', usernames)
        self.assertNotIn('bob', usernames)

    def test_no_search_returns_all_users(self):
        """Export without search includes all visible users."""
        UserSkill.objects.create(
            user=self.user1, skill=self.skill1, level=3,
            last_updated_by=self.user1,
        )
        UserSkill.objects.create(
            user=self.user2, skill=self.skill1, level=4,
            last_updated_by=self.user2,
        )
        resp = export_matrix_csv({self.user1.id, self.user2.id})
        rows = self._parse_csv(resp)
        usernames = [r[0] for r in rows[1:]]
        self.assertIn('alice', usernames)
        self.assertIn('bob', usernames)


class ExportMatrixXlsxTest(TestCase):
    """The wide XLSX export: one row per person, one column per skill.

    Deliberately a different shape from ``export_matrix_csv``, which stays the
    long one-row-per-rating format. Both remain reachable.
    """

    def setUp(self):
        self.backend = SkillCategory.objects.create(name="Backend", code="backend")
        self.infra = SkillCategory.objects.create(name="Infra", code="infra")
        self.python = Skill.objects.create(
            category=self.backend, name="Python", code="python"
        )
        self.django = Skill.objects.create(
            category=self.backend, name="Django", code="django"
        )
        self.linux = Skill.objects.create(
            category=self.infra, name="Linux", code="linux"
        )
        self.alice = User.objects.create_user(
            username="alice", password="x", first_name="Alice", last_name="Xhelili"
        )
        self.bob = User.objects.create_user(
            username="bob", password="x", first_name="Bob", last_name="Bitri"
        )
        UserSkill.objects.create(
            user=self.alice, skill=self.python, level=4, last_updated_by=self.alice
        )
        UserSkill.objects.create(
            user=self.alice, skill=self.linux, level=2, last_updated_by=self.alice
        )

    def _sheet(self, response, name='Skill matrix'):
        return load_workbook(io.BytesIO(response.content))[name]

    def test_response_headers(self):
        resp = export_matrix_xlsx({self.alice.id})
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment', resp['Content-Disposition'])
        self.assertIn('skills_matrix.xlsx', resp['Content-Disposition'])

    def test_two_header_rows_band_then_skill(self):
        ws = self._sheet(export_matrix_xlsx({self.alice.id}))
        band = [c.value for c in ws[1]]
        header = [c.value for c in ws[2]]
        self.assertEqual(header[:4], ['Surname', 'Name', 'Username', 'Grade'])
        # Skill columns follow, grouped by category, with the category name on
        # row 1 over the first column of its group only.
        self.assertEqual(header[4:], ['Django', 'Python', 'Linux'])
        self.assertEqual(band[4], 'Backend')
        self.assertIsNone(band[5])
        self.assertEqual(band[6], 'Infra')

    def test_one_row_per_person_with_level_names(self):
        ws = self._sheet(export_matrix_xlsx({self.alice.id, self.bob.id}))
        rows = {r[2]: r for r in ws.iter_rows(min_row=3, values_only=True)}
        self.assertEqual(set(rows), {'alice', 'bob'})
        self.assertEqual(rows['alice'][0], 'Xhelili')
        self.assertEqual(rows['alice'][1], 'Alice')
        # Django unrated -> blank, Python L4, Linux L2.
        self.assertEqual(rows['alice'][4:], (None, 'Advanced', 'Developing'))
        # An unrated person still gets a row, so gaps stay visible.
        self.assertEqual(rows['bob'][4:], (None, None, None))

    def test_uses_admin_renamed_level_labels(self):
        labels = SkillLevelLabels.get_singleton()
        labels.level_4_label = 'Expert'
        labels.save(update_fields=['level_4_label'])
        ws = self._sheet(export_matrix_xlsx({self.alice.id}))
        row = next(ws.iter_rows(min_row=3, values_only=True))
        self.assertEqual(row[5], 'Expert')

    def test_lists_sheet_carries_the_scale(self):
        wb = load_workbook(io.BytesIO(export_matrix_xlsx({self.alice.id}).content))
        self.assertIn('Lists', wb.sheetnames)
        values = [r[0] for r in wb['Lists'].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(
            values, ['Foundational', 'Developing', 'Proficient', 'Advanced', 'Mastery']
        )

    def test_grade_column_carries_tech_and_level(self):
        tech = Tech.objects.create(name="Infrastructure", code="INFRA")
        level = TechLevel.objects.create(tech=tech, name="L3", code="L3", rank=3)
        UserTech.objects.create(
            user_profile=self.alice.profile, tech=tech, level=level
        )
        ws = self._sheet(export_matrix_xlsx({self.alice.id}))
        row = next(ws.iter_rows(min_row=3, values_only=True))
        self.assertEqual(row[3], 'Infrastructure L3')

    def test_category_filter_narrows_the_columns(self):
        ws = self._sheet(export_matrix_xlsx({self.alice.id}, category_code='backend'))
        header = [c.value for c in ws[2]]
        self.assertEqual(header[4:], ['Django', 'Python'])

    def test_search_narrows_the_rows(self):
        ws = self._sheet(
            export_matrix_xlsx({self.alice.id, self.bob.id}, search='alice')
        )
        usernames = [r[2] for r in ws.iter_rows(min_row=3, values_only=True)]
        self.assertEqual(usernames, ['alice'])

    def test_inactive_users_are_excluded(self):
        self.bob.is_active = False
        self.bob.save(update_fields=['is_active'])
        ws = self._sheet(export_matrix_xlsx({self.alice.id, self.bob.id}))
        usernames = [r[2] for r in ws.iter_rows(min_row=3, values_only=True)]
        self.assertEqual(usernames, ['alice'])

    def test_no_visible_users_still_returns_a_valid_workbook(self):
        ws = self._sheet(export_matrix_xlsx(set()))
        self.assertEqual(ws.max_row, 2)

    def test_query_count_is_flat_in_the_number_of_people(self):
        """One pass over the ratings, not one query per person."""
        # Warm the label singleton: its first ``get_or_create`` writes a row and
        # would make the first measurement the expensive one.
        SkillLevelLabels.get_singleton()
        with CaptureQueriesContext(connection) as few:
            export_matrix_xlsx({self.alice.id, self.bob.id})
        extra = []
        for i in range(10):
            user = User.objects.create_user(username=f"qb{i}", password="x")
            UserSkill.objects.create(
                user=user, skill=self.python, level=3, last_updated_by=user
            )
            extra.append(user.id)
        with CaptureQueriesContext(connection) as many:
            export_matrix_xlsx({self.alice.id, self.bob.id, *extra})
        self.assertEqual(
            len(few.captured_queries),
            len(many.captured_queries),
            f"queries grew {len(few.captured_queries)} -> "
            f"{len(many.captured_queries)} as people went 2 -> 12",
        )
