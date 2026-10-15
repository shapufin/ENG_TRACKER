"""TL scorecard summary export: XLSX (default) and PDF branches of one endpoint."""
import io
from datetime import date

from django.core.management import call_command
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook
from pypdf import PdfReader
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.permissions.models import Role, UserRole
from core.utils.csv_injection_assert import PAYLOAD, assert_xlsx_neutralised

from .excel_export import build_workbook_bytes
from .report_kpis import kpi_rows, status_for
from .models import Absence, PIPRecord
from .services import build_scorecard
from .test_tl_scorecard_hbpr import HbprScorecardBase
from .testing import make_user as _make_user
from .viewsets import TLScorecardViewSet

XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
SHEETS = ['Summary', 'Trend', 'Governance', 'HBPR partnership log']


def _assign_tl_role(user, code='italian_tl'):
    role, _ = Role.objects.get_or_create(code=code, defaults={'name': code})
    UserRole.objects.get_or_create(user=user, role=role, defaults={'is_active': True})


def _pdf_text(content):
    return ' '.join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)


def _find_row(ws, label, column=2):
    for row in ws.iter_rows(min_row=1):
        if row[column - 1].value == label:
            return row
    raise AssertionError(f'{label!r} not found in {ws.title}')


class SummaryExportFormatTests(TestCase):
    def setUp(self):
        self.factory = APIRequestFactory()
        call_command('seed_plugin_permissions')
        self.leader = _make_user('leader_formats')
        _assign_tl_role(self.leader)
        self.member = _make_user('member_formats')
        self.member.profile.italian_tl = self.leader
        self.member.profile.save()
        PIPRecord.objects.create(employee=self.member, tl=self.leader, start_date=date.today())
        Absence.objects.create(employee=self.member, flagged_by=self.leader, absence_date=date.today())

    def _get(self, params=None):
        request = self.factory.get('/api/plugins/tl_scorecard/export/', params or {})
        force_authenticate(request, user=self.leader)
        return TLScorecardViewSet.as_view({'get': 'export'})(request)

    def test_pdf_branch_returns_a_pdf_attachment(self):
        resp = self._get({'file_format': 'pdf'})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF'))
        month = date.today().strftime('%Y-%m')
        self.assertIn(
            f'attachment; filename="tl_scorecard_leader_formats_{month}.pdf"',
            resp['Content-Disposition'],
        )

    def test_xlsx_is_the_default_and_the_explicit_choice(self):
        for params in ({}, {'file_format': 'xlsx'}):
            with self.subTest(params=params):
                resp = self._get(params)
                self.assertEqual(resp.status_code, 200)
                self.assertEqual(resp['Content-Type'], XLSX_TYPE)
                self.assertEqual(resp.content[:2], b'PK')
                self.assertTrue(resp['Content-Disposition'].endswith('.xlsx"'))

    def test_workbook_sheets_and_numeric_one_to_one_summary(self):
        wb = load_workbook(io.BytesIO(self._get({'file_format': 'xlsx'}).content))
        self.assertEqual(wb.sheetnames, SHEETS)
        row = _find_row(wb['Summary'], '1:1 compliance')
        self.assertIsInstance(row[2].value, (int, float))
        self.assertNotIsInstance(row[2].value, bool)

    def test_unknown_file_format_is_rejected(self):
        resp = self._get({'file_format': 'docx'})
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data, {'error': 'file_format must be xlsx or pdf'})


class WorkbookFormulaSafetyTests(TestCase):
    """User-controlled text is stored as a literal, never a formula or link."""

    def _build(self, **evidence_fields):
        leader = _make_user('leader_formula_safety')
        evidence = [{
            'kind_display': 'Cadence', 'occurred_on': '2026-10-01', 'reporting_year': 2026,
            'shared_summary': '', 'action_items': '', 'reference_url': '', **evidence_fields,
        }]
        governance = {
            'open_pips': [], 'pending_promotions': [], 'epr_cycles': [],
            'open_absences': [{'employee': PAYLOAD, 'absence_date': '2026-10-01', 'reason': PAYLOAD}],
        }
        return build_workbook_bytes(
            build_scorecard(leader, date.today()), governance, 'Oct 2026', hbpr_evidence=evidence,
        )

    def test_partnership_log_shared_summary_is_neutralised(self):
        assert_xlsx_neutralised(self, self._build(shared_summary=PAYLOAD))

    def test_partnership_log_action_items_is_neutralised(self):
        assert_xlsx_neutralised(self, self._build(action_items=PAYLOAD))

    def test_partnership_log_reference_url_is_neutralised(self):
        assert_xlsx_neutralised(self, self._build(reference_url=PAYLOAD))

    def test_tab_leading_text_is_neutralised(self):
        wb = load_workbook(io.BytesIO(self._build(shared_summary='	=cmd')))
        values = [c for row in wb['HBPR partnership log'].iter_rows(values_only=True) for c in row]
        self.assertIn("'	=cmd", values)

    def test_governance_free_text_is_neutralised(self):
        assert_xlsx_neutralised(self, self._build())


class HbprExportFormatTests(HbprScorecardBase):
    """HBPR sees governance, never the leader's employee one-on-ones."""

    def setUp(self):
        super().setUp()
        from .models import Meeting

        second = _make_user('second_member_formats')
        second.profile.albanian_tl = self.tl
        second.profile.save()
        # One of two direct reports met 1:1 this month -> a real, non-trivial %.
        Meeting.objects.create(
            organizer=self.tl, counterparty=self.member, meeting_type='one_on_one',
            occurred_on=date.today(),
        )
        self.expected_pct = build_scorecard(self.tl, date.today())['meetings']['one_on_one_compliance_pct']
        self.assertEqual(self.expected_pct, 50.0)

    def _export(self, file_format):
        request = self.factory.get('/x/', {'leader_id': self.tl.id, 'file_format': file_format})
        force_authenticate(request, user=self.hbpr)
        return TLScorecardViewSet.as_view({'get': 'export'})(request)

    def test_xlsx_one_to_one_is_a_dash_in_summary_and_every_trend_row(self):
        resp = self._export('xlsx')
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(_find_row(wb['Summary'], '1:1 compliance')[2].value, '—')

        trend = wb['Trend']
        header = next(row for row in trend.iter_rows(values_only=True) if '1:1 compliance' in row)
        column = header.index('1:1 compliance')
        data_rows = [
            row for row in trend.iter_rows(min_row=2, values_only=True)
            if row[0] is not None
        ]
        self.assertEqual(len(data_rows), 6)
        for row in data_rows:
            self.assertEqual(row[column], '—')

    def test_pdf_never_contains_the_leaders_one_to_one_percentage(self):
        resp = self._export('pdf')
        self.assertEqual(resp.status_code, 200)
        text = _pdf_text(resp.content)
        self.assertIn('Meetings & reviews', text)
        self.assertNotIn(f'{self.expected_pct}%', text)

    def test_governance_is_limited_to_hbpr_scoped_subjects(self):
        ex_member = _make_user('ex_member_formats', first_name='Former', last_name='Colleague')
        PIPRecord.objects.create(employee=ex_member, tl=self.tl, start_date=date.today())
        PIPRecord.objects.create(employee=self.member, tl=self.tl, start_date=date.today())

        wb = load_workbook(io.BytesIO(self._export('xlsx').content))
        governance_text = ' '.join(
            str(cell) for row in wb['Governance'].iter_rows(values_only=True) for cell in row if cell
        )
        self.assertNotIn('ex_member_formats', governance_text)
        self.assertIn(self.member.username, governance_text)

        pdf_text = _pdf_text(self._export('pdf').content)
        self.assertNotIn('ex_member_formats', pdf_text)
        self.assertIn(self.member.username, pdf_text)


class ReportKpiStatusTests(SimpleTestCase):
    def test_higher_is_better_targets(self):
        self.assertEqual(status_for(100.0, 100.0), 'On track')
        self.assertEqual(status_for(85.0, 100.0), 'Watch')
        self.assertEqual(status_for(50.0, 100.0), 'Off track')
        self.assertEqual(status_for(None, 100.0), '—')

    def test_zero_targets_follow_the_ui_tone_of_a_breach(self):
        self.assertEqual(status_for(0, 0, nonzero='Watch'), 'On track')
        self.assertEqual(status_for(1, 0, nonzero='Watch'), 'Watch')
        self.assertEqual(status_for(2, 0, nonzero='Off track'), 'Off track')

    def test_review_target_is_pro_rata_to_the_month(self):
        rows = kpi_rows({
            'month': '2026-06-01', 'team_size': 1,
            'review_deliveries_ytd': 6,
            'leave': {'pct_within_2_days': None, 'pending_at_month_end': 0},
            'overtime': {'avg_turnaround_days': None},
            'meetings': {'one_on_one_compliance_pct': None, 'tl_sync_count': 0,
                         'team_meetings_with_hrbp': 0},
            'idle': {'open_count': 0}, 'absences': {'breached_5_day_sla': 0},
            'pip': {'pending_approval_count': 0}, 'promotion': {'promoted_pct': None},
            'escalation_count': 0,
        })
        reviews = next(r for r in rows if r['metric'].startswith('Management reviews'))
        self.assertEqual(reviews['target'], 6)
        self.assertEqual(reviews['status'], 'On track')
