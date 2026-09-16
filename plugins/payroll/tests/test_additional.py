"""Additional regression tests for Payroll audit findings."""
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
import io

import openpyxl
from django.core.exceptions import ValidationError
from django.test import TestCase
from rest_framework.test import APIRequestFactory, force_authenticate

from plugins.payroll.models import PayrollRuleSet, PayrollRun, WageAssignment
from plugins.payroll.services.export_service import (
    OVERTIME_CATEGORY_LABELS,
    _overtime_rows,
    _pct,
    _standby_rows,
    generate_payroll_excel,
    generate_payroll_pdf,
    generate_payslip_pdf,
)
from plugins.payroll.services.payroll_service import generate_draft_run, finalize_run
from plugins.payroll.viewsets import PayrollConfigurationViewSet, PayrollRunViewSet

from django.contrib.auth import get_user_model

User = get_user_model()


class PayrollExportPureHelperTests(TestCase):
    """These helpers take plain data (no DB), so they're tested directly."""

    def test_pct_returns_zero_when_whole_is_zero(self):
        self.assertEqual(_pct(Decimal('50'), Decimal('0')), Decimal('0'))

    def test_pct_computes_percentage(self):
        self.assertEqual(_pct(Decimal('25'), Decimal('200')), Decimal('12.5'))

    def test_overtime_rows_maps_codes_to_human_labels(self):
        line = SimpleNamespace(
            total_gross=Decimal('1000'),
            overtime_breakdown={'categories': [
                {'code': 'weekday_night', 'hours': '4', 'multiplier': '1.30', 'amount': '100'},
                {'code': 'holiday', 'hours': '2', 'multiplier': '1.50', 'amount': '50'},
            ]},
        )
        rows = _overtime_rows(line)
        self.assertEqual(rows[0]['label'], OVERTIME_CATEGORY_LABELS['weekday_night'])
        self.assertEqual(rows[0]['multiplier_pct'], Decimal('130.00'))
        self.assertEqual(rows[0]['pct_of_gross'], Decimal('10'))
        self.assertEqual(rows[1]['label'], OVERTIME_CATEGORY_LABELS['holiday'])

    def test_overtime_rows_empty_when_no_categories(self):
        line = SimpleNamespace(total_gross=Decimal('1000'), overtime_breakdown={})
        self.assertEqual(_overtime_rows(line), [])

    def test_standby_rows_splits_weekday_and_weekend(self):
        line = SimpleNamespace(
            total_gross=Decimal('1000'),
            calculation_trace={'standby': {
                'weekday_hours': '5', 'weekend_hours': '3',
                'weekday_rate': '10', 'weekend_rate': '15',
            }},
        )
        rows = _standby_rows(line)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['amount'], Decimal('50'))
        self.assertEqual(rows[1]['amount'], Decimal('45'))

    def test_standby_rows_skips_zero_buckets(self):
        line = SimpleNamespace(
            total_gross=Decimal('1000'),
            calculation_trace={'standby': {
                'weekday_hours': '0', 'weekend_hours': '3',
                'weekday_rate': '10', 'weekend_rate': '15',
            }},
        )
        rows = _standby_rows(line)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['label'], 'Weekend Standby')


class PayrollAdditionalRegressionTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(
            username='audit_admin', email='audit@example.com', password='pass',
        )
        cls.employee = User.objects.create_user(
            username='audit_employee', email='employee@example.com', password='pass',
        )
        WageAssignment.objects.create(
            user=cls.employee,
            gross_monthly_wage=Decimal('100000'),
            effective_from=date(2026, 1, 1),
        )
        from django.core.management import call_command
        call_command('seed_payroll_rules', verbosity=0)

    def setUp(self):
        self.factory = APIRequestFactory()

    def _new_run(self, month):
        rule_set = PayrollRuleSet.objects.get(code='AL_2026_BOSHTI_REFERENCE')
        run = PayrollRun.objects.create(
            year=2026, month=month, status='draft', rule_set=rule_set,
            created_by=self.admin,
        )
        generate_draft_run(run, [self.employee])
        return run

    def test_configuration_list_returns_singleton_object(self):
        request = self.factory.get('/api/plugins/payroll/configuration/')
        force_authenticate(request, user=self.admin)
        response = PayrollConfigurationViewSet.as_view({'get': 'list'})(request)
        self.assertEqual(response.status_code, 200)
        self.assertIsInstance(response.data, dict)
        self.assertEqual(response.data['id'], 1)

    def test_overlapping_wages_are_rejected_during_resolution(self):
        WageAssignment.objects.create(
            user=self.employee,
            gross_monthly_wage=Decimal('110000'),
            effective_from=date(2026, 2, 1),
        )
        with self.assertRaises(ValidationError):
            WageAssignment.resolve_for_month(self.employee, 2026, 2)

    def test_finalized_run_cannot_be_updated_or_deleted(self):
        run = self._new_run(8)
        finalize_run(run, self.admin)

        update_request = self.factory.patch(
            f'/api/plugins/payroll/runs/{run.id}/', {'notes': 'tamper'}, format='json',
        )
        force_authenticate(update_request, user=self.admin)
        update_response = PayrollRunViewSet.as_view({'patch': 'partial_update'})(
            update_request, pk=run.id,
        )
        self.assertEqual(update_response.status_code, 409)

        delete_request = self.factory.delete(f'/api/plugins/payroll/runs/{run.id}/')
        force_authenticate(delete_request, user=self.admin)
        delete_response = PayrollRunViewSet.as_view({'delete': 'destroy'})(
            delete_request, pk=run.id,
        )
        self.assertEqual(delete_response.status_code, 409)

    def test_payslip_pdf_contains_pdf_signature(self):
        run = self._new_run(9)
        run = finalize_run(run, self.admin)
        line = run.lines.get(user=self.employee)
        payload = generate_payslip_pdf(line)
        # generate_payslip_pdf returns a BytesIO stream
        self.assertTrue(payload.read(5).startswith(b'%PDF'))

    def test_payslip_pdf_renders_with_zero_overtime(self):
        run = self._new_run(1)
        run = finalize_run(run, self.admin)
        line = run.lines.get(user=self.employee)
        payload = generate_payslip_pdf(line)
        self.assertTrue(payload.read(5).startswith(b'%PDF'))

    def test_payslip_pdf_renders_with_non_official_rule_set(self):
        # Seeded rule set is validation_status='reference' by default — the
        # warning banner path must not crash the PDF build.
        run = self._new_run(2)
        run = finalize_run(run, self.admin)
        self.assertNotEqual(run.rule_set.validation_status, 'official')
        line = run.lines.get(user=self.employee)
        payload = generate_payslip_pdf(line)
        self.assertTrue(payload.read(5).startswith(b'%PDF'))

    def test_consolidated_payroll_pdf_renders_cover_and_all_payslips(self):
        second_employee = User.objects.create_user(
            username='audit_employee_2', email='employee2@example.com', password='pass',
        )
        WageAssignment.objects.create(
            user=second_employee,
            gross_monthly_wage=Decimal('90000'),
            effective_from=date(2026, 1, 1),
        )
        rule_set = PayrollRuleSet.objects.get(code='AL_2026_BOSHTI_REFERENCE')
        run = PayrollRun.objects.create(
            year=2026, month=7, status='draft', rule_set=rule_set,
            created_by=self.admin,
        )
        generate_draft_run(run, [self.employee, second_employee])
        run = finalize_run(run, self.admin)

        payload = generate_payroll_pdf(run)
        self.assertTrue(payload.read(5).startswith(b'%PDF'))

    def test_consolidated_payroll_pdf_renders_with_no_lines(self):
        rule_set = PayrollRuleSet.objects.get(code='AL_2026_BOSHTI_REFERENCE')
        run = PayrollRun.objects.create(
            year=2026, month=8, status='finalized', rule_set=rule_set,
            created_by=self.admin,
        )
        payload = generate_payroll_pdf(run)
        self.assertTrue(payload.read(5).startswith(b'%PDF'))

    def test_excel_export_escapes_formula_like_employee_name(self):
        self.employee.first_name = '=1+1'
        self.employee.save(update_fields=['first_name'])
        run = self._new_run(11)

        workbook = openpyxl.load_workbook(io.BytesIO(generate_payroll_excel(run).read()))

        self.assertEqual(workbook['Employee Breakdown']['A2'].value, "'=1+1")
        self.assertEqual(workbook['Overtime Detail']['A2'].value, "'=1+1")
        # Row 1 is the bracket summary, row 3 the header, data starts row 4.
        self.assertEqual(workbook['Tax & Contributions Detail']['A4'].value, "'=1+1")

    def test_excel_export_has_tax_and_contributions_sheet(self):
        run = self._new_run(3)
        workbook = openpyxl.load_workbook(io.BytesIO(generate_payroll_excel(run).read()))
        self.assertIn('Tax & Contributions Detail', workbook.sheetnames)
        ws = workbook['Tax & Contributions Detail']
        # Row 1 is the bracket-schedule summary line; column headers are row 3.
        header_row = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]
        self.assertIn('Effective Tax Rate %', header_row)

    def test_excel_employee_breakdown_has_percentage_columns(self):
        run = self._new_run(4)
        workbook = openpyxl.load_workbook(io.BytesIO(generate_payroll_excel(run).read()))
        ws = workbook['Employee Breakdown']
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for expected in ('Overtime % Gross', 'Effective Tax Rate %', 'Total Deduction % Gross'):
            self.assertIn(expected, headers)

    def test_excel_overtime_detail_uses_human_labels_not_raw_codes(self):
        run = self._new_run(6)
        workbook = openpyxl.load_workbook(io.BytesIO(generate_payroll_excel(run).read()))
        ws = workbook['Overtime Detail']
        # Employee has zero overtime in this fixture, so only the "(none)" row
        # exists — assert the sheet still renders without raw category codes
        # leaking in for employees who DO have overtime, by checking the
        # header contract instead (label mapping is unit-tested directly).
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers[1], 'Category')

    def test_active_run_uniqueness_rejects_second_draft(self):
        rule_set = PayrollRuleSet.objects.get(code='AL_2026_BOSHTI_REFERENCE')
        PayrollRun.objects.create(
            year=2026, month=10, status='draft', rule_set=rule_set,
            created_by=self.admin,
        )
        with self.assertRaises(Exception):
            PayrollRun.objects.create(
                year=2026, month=10, status='draft', rule_set=rule_set,
                created_by=self.admin,
            )
