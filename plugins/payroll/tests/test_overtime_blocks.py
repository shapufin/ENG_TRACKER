"""Overtime paid blocks in payroll: freezing, stale drafts and the export rows."""
import io
from datetime import date, time
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.test import TestCase
from pypdf import PdfReader

from apps.overtime.models import Client, OvertimeLog
from plugins.payroll.models import PayrollConfiguration, PayrollRuleSet, PayrollRun, WageAssignment
from plugins.payroll.services.export_service import (
    generate_payroll_excel,
    generate_payroll_pdf,
    generate_payslip_pdf,
)
from plugins.payroll.services.payroll_service import finalize_run, generate_draft_run

User = get_user_model()
WORK_DAY = date(2026, 9, 14)
PERIOD = date(2026, 9, 1)


class OvertimeBlockPayrollTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        PayrollConfiguration.objects.update_or_create(pk=1, defaults={'require_tl_closed_before_finalize': False})
        call_command('seed_payroll_rules', verbosity=0)
        cls.admin = User.objects.create_superuser(username='blk_admin', email='a@example.com', password='x')
        cls.employee = User.objects.create_user(
            username='blk_emp', email='e@example.com', password='x', first_name='Block', last_name='Worker',
        )
        cls.client_obj = Client.objects.create(name='Block Client', code='BLK')
        WageAssignment.objects.create(
            user=cls.employee, gross_monthly_wage=Decimal('100000'), effective_from=date(2026, 1, 1),
        )

    def _log(self, start, end, evidence='', status='approved'):
        return OvertimeLog.objects.create(
            user=self.employee, client=self.client_obj, date=WORK_DAY, hours=1,
            start_time=start, end_time=end, evidence_type='ticket', evidence=evidence,
            requested_processing_period=PERIOD, status=status,
        )

    def _run(self):
        run = PayrollRun.objects.create(
            year=2026, month=9, status='draft',
            rule_set=PayrollRuleSet.objects.get(code='AL_2026_BOSHTI_REFERENCE'),
            created_by=self.admin,
        )
        generate_draft_run(run, [self.employee])
        return run

    def test_a_split_hour_is_paid_once(self):
        self._log(time(8, 0), time(8, 30), 'INC-1')
        self._log(time(8, 31), time(9, 0), 'INC-2')
        line = self._run().lines.get(user=self.employee)
        self.assertEqual(line.overtime_hours, Decimal('1.00'))

    def test_an_entry_added_after_a_paid_one_gets_only_what_is_left(self):
        first = self._log(time(8, 0), time(8, 30), 'INC-1')
        run = self._run()
        finalize_run(run, self.admin)
        second = self._log(time(8, 31), time(9, 0), 'INC-2')
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual((first.hours, second.hours), (Decimal('1.00'), Decimal('0.00')))
        second_run = PayrollRun.objects.create(
            year=2026, month=10, status='draft',
            rule_set=PayrollRuleSet.objects.get(code='AL_2026_BOSHTI_REFERENCE'),
            created_by=self.admin,
        )
        generate_draft_run(second_run, [self.employee])
        self.assertEqual(second_run.lines.get(user=self.employee).overtime_hours, Decimal('0.00'))

    def test_a_draft_cannot_be_finalized_after_its_overtime_shares_changed(self):
        self._log(time(8, 0), time(8, 20), 'INC-1')
        self._log(time(8, 50), time(9, 10), 'INC-2')
        run = self._run()
        self._log(time(8, 30), time(8, 40), 'INC-3')
        with self.assertRaisesMessage(ValidationError, 'Regenerate'):
            finalize_run(run, self.admin)

    def _finalized_run(self):
        self._log(time(8, 0), time(8, 30), 'INC-1')
        self._log(time(8, 31), time(9, 0), 'INC-2')
        self._log(time(20, 0), time(20, 30), 'CALL-7')
        run = self._run()
        finalize_run(run, self.admin)
        return run

    def test_excel_lists_one_row_per_block_with_ranges_and_references(self):
        run = self._finalized_run()
        workbook = openpyxl.load_workbook(io.BytesIO(generate_payroll_excel(run).read()))
        rows = list(workbook['Overtime Entries'].iter_rows(values_only=True))
        self.assertEqual(
            rows[0],
            ('Employee', 'Date', 'Paid Window', 'Entries (time, reference)', 'Worked (min)', 'Paid Hours'),
        )
        first, second = rows[1], rows[2]
        self.assertEqual(first[:3], ('Block Worker', '2026-09-14', '08:00-09:00'))
        self.assertEqual(first[3], '08:00-08:30 INC-1\n08:31-09:00 INC-2')
        self.assertEqual((first[4], Decimal(str(first[5]))), (59, Decimal('1')))
        self.assertEqual(second[2:4], ('20:00-21:00', '20:00-20:30 CALL-7'))

    def test_excel_text_cells_are_formula_safe(self):
        self._log(time(8, 0), time(8, 30), '=HYPERLINK("http://x")')
        run = self._run()
        finalize_run(run, self.admin)
        workbook = openpyxl.load_workbook(io.BytesIO(generate_payroll_excel(run).read()))
        cell = workbook['Overtime Entries'].cell(row=2, column=4).value
        self.assertNotIn(cell[0], '=+-@\t\r')

    def test_payslip_and_consolidated_pdf_show_the_overtime_entries(self):
        run = self._finalized_run()
        slip = PdfReader(io.BytesIO(generate_payslip_pdf(run.lines.get(user=self.employee)).read()))
        text = ''.join(page.extract_text() for page in slip.pages)
        self.assertIn('Overtime Entries', text)
        self.assertIn('INC-1', text)
        self.assertIn('08:00', text)
        everyone = PdfReader(io.BytesIO(generate_payroll_pdf(run).read()))
        self.assertIn('INC-2', ''.join(page.extract_text() for page in everyone.pages))
